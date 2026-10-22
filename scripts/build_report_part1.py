"""
The Cloud Scraper — Serverless Automated Data Pipeline
Course Project Final Report Generator
Course: BCSE355L - Cloud Architecture and Design
"""

import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from report_helpers import (
    COLOR_NAVY, COLOR_BLUE, COLOR_SLATE, COLOR_DARK, COLOR_MUTED,
    HEX_NAVY, HEX_BLUE, HEX_LIGHT_BG, HEX_ALT_ROW, HEX_BORDER,
    set_cell_background, set_cell_margins, set_table_borders,
    format_table_header_and_rows, add_custom_heading, add_body_paragraph,
    add_bullet_item, add_code_block, add_callout, add_figure_with_caption,
    add_page_number_to_run, add_numpages_to_run, setup_document_styles_and_margins
)

PROJECT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUTPUT_DOCX = os.path.join(PROJECT_DIR, "The_Cloud_Scraper_Final_Project_Report.docx")
FIGURES_DIR = os.path.join(PROJECT_DIR, "figures")
IMAGES_DIR = os.path.join(PROJECT_DIR, "images")


def create_cover_page(doc):
    """Generate the official academic cover page matching Section 1 specifications exactly."""
    # Ensure no header/footer on cover page
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.space_before = Pt(36)
    p_inst.paragraph_format.space_after = Pt(4)
    run_inst = p_inst.add_run("VELLORE INSTITUTE OF TECHNOLOGY")
    run_inst.font.name = "Calibri"
    run_inst.font.size = Pt(14)
    run_inst.font.bold = True
    run_inst.font.color.rgb = COLOR_NAVY

    p_school = doc.add_paragraph()
    p_school.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_school.paragraph_format.space_before = Pt(0)
    p_school.paragraph_format.space_after = Pt(2)
    run_school = p_school.add_run("School of Computer Science Engineering and Information Systems")
    run_school.font.name = "Calibri"
    run_school.font.size = Pt(12)
    run_school.font.bold = True
    run_school.font.color.rgb = COLOR_SLATE

    p_prog = doc.add_paragraph()
    p_prog.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_prog.paragraph_format.space_before = Pt(0)
    p_prog.paragraph_format.space_after = Pt(40)
    run_prog = p_prog.add_run("Programme: B.Tech Information Technology")
    run_prog.font.name = "Calibri"
    run_prog.font.size = Pt(11)
    run_prog.font.color.rgb = COLOR_SLATE

    # Project Title Box
    p_type = doc.add_paragraph()
    p_type.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_type.paragraph_format.space_before = Pt(0)
    p_type.paragraph_format.space_after = Pt(6)
    run_type = p_type.add_run("COURSE PROJECT FINAL REPORT")
    run_type.font.name = "Calibri"
    run_type.font.size = Pt(11)
    run_type.font.bold = True
    run_type.font.color.rgb = COLOR_BLUE

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(8)
    run_title = p_title.add_run("The Cloud Scraper — Serverless Automated Data Pipeline")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_NAVY

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(45)
    run_sub = p_sub.add_run("A General-Purpose, Event-Driven Web Harvesting and ETL Architecture on AWS")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True
    run_sub.font.color.rgb = COLOR_SLATE

    # Course Metadata Table
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Subject:", "Cloud Architecture and Design"),
        ("Course Code:", "BCSE355L"),
        ("Semester:", "Fall 2026–2027"),
        ("Slot:", "B1 + TB1")
    ]
    for row_idx, (k, v) in enumerate(meta_data):
        row = meta_table.rows[row_idx]
        cell_k, cell_v = row.cells[0], row.cells[1]
        cell_k.width = Inches(2.2)
        cell_v.width = Inches(3.8)
        set_cell_background(cell_k, "FFFFFF")
        set_cell_background(cell_v, "FFFFFF")
        set_cell_margins(cell_k, top=40, bottom=40, left=40, right=40)
        set_cell_margins(cell_v, top=40, bottom=40, left=40, right=40)
        
        pk = cell_k.paragraphs[0]
        pk.paragraph_format.space_before = Pt(0)
        pk.paragraph_format.space_after = Pt(0)
        rk = pk.add_run(k)
        rk.font.name = "Calibri"
        rk.font.bold = True
        rk.font.size = Pt(10.5)
        rk.font.color.rgb = COLOR_NAVY

        pv = cell_v.paragraphs[0]
        pv.paragraph_format.space_before = Pt(0)
        pv.paragraph_format.space_after = Pt(0)
        rv = pv.add_run(v)
        rv.font.name = "Calibri"
        rv.font.size = Pt(10.5)
        rv.font.color.rgb = COLOR_DARK

    set_table_borders(meta_table, "FFFFFF")

    # Spacing
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(30)
    p_space.paragraph_format.space_after = Pt(0)

    # Submission Details (Centred)
    p_subto = doc.add_paragraph()
    p_subto.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_subto.paragraph_format.space_before = Pt(0)
    p_subto.paragraph_format.space_after = Pt(2)
    run_subto = p_subto.add_run("Submitted To:")
    run_subto.font.name = "Calibri"
    run_subto.font.size = Pt(10.5)
    run_subto.font.bold = True
    run_subto.font.color.rgb = COLOR_SLATE

    p_prof = doc.add_paragraph()
    p_prof.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_prof.paragraph_format.space_before = Pt(0)
    p_prof.paragraph_format.space_after = Pt(20)
    run_prof = p_prof.add_run("Prof. Venkatesh P")
    run_prof.font.name = "Calibri"
    run_prof.font.size = Pt(13)
    run_prof.font.bold = True
    run_prof.font.color.rgb = COLOR_NAVY

    p_by = doc.add_paragraph()
    p_by.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_by.paragraph_format.space_before = Pt(0)
    p_by.paragraph_format.space_after = Pt(4)
    run_by = p_by.add_run("Submitted By:")
    run_by.font.name = "Calibri"
    run_by.font.size = Pt(10.5)
    run_by.font.bold = True
    run_by.font.color.rgb = COLOR_SLATE

    team_members = [
        "1. Bhumi Choudhary — 24BIT0586",
        "2. Kaushal More — 24BIT0283",
        "3. Kshitij Agarwal — 24BIT0317"
    ]
    for mem in team_members:
        p_mem = doc.add_paragraph()
        p_mem.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_mem.paragraph_format.space_before = Pt(0)
        p_mem.paragraph_format.space_after = Pt(2)
        run_mem = p_mem.add_run(mem)
        run_mem.font.name = "Calibri"
        run_mem.font.size = Pt(11)
        run_mem.font.bold = True
        run_mem.font.color.rgb = COLOR_DARK

    doc.add_page_break()


def configure_header_footer(doc):
    """Set up running header and footer for all pages after the cover page."""
    section = doc.sections[0]
    
    # Header
    header = section.header
    p_head = header.paragraphs[0]
    p_head.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_head.paragraph_format.space_after = Pt(0)
    r_head = p_head.add_run("The Cloud Scraper — Serverless Automated Data Pipeline | BCSE355L")
    r_head.font.name = "Calibri"
    r_head.font.size = Pt(8.5)
    r_head.font.color.rgb = COLOR_MUTED

    # Footer
    footer = section.footer
    p_foot = footer.paragraphs[0]
    p_foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_foot.paragraph_format.space_after = Pt(0)
    
    r_foot_left = p_foot.add_run("School of Computer Science Engineering and Information Systems (VIT)          ")
    r_foot_left.font.name = "Calibri"
    r_foot_left.font.size = Pt(8.5)
    r_foot_left.font.color.rgb = COLOR_MUTED

    r_foot_pg = p_foot.add_run("Page ")
    r_foot_pg.font.name = "Calibri"
    r_foot_pg.font.size = Pt(8.5)
    r_foot_pg.font.color.rgb = COLOR_SLATE
    add_page_number_to_run(r_foot_pg)

    r_foot_of = p_foot.add_run(" of ")
    r_foot_of.font.name = "Calibri"
    r_foot_of.font.size = Pt(8.5)
    r_foot_of.font.color.rgb = COLOR_SLATE
    add_numpages_to_run(r_foot_of)


def create_abstract_and_acknowledgements(doc):
    """Create the Abstract and Acknowledgement sections."""
    add_custom_heading(doc, "Abstract", level=1)
    
    add_body_paragraph(
        doc,
        "Modern computational organizations increasingly depend on timely, high-fidelity structured data harvested "
        "from dynamic web sources such as e-commerce catalogs, news aggregators, research repositories, and public indices. "
        "However, traditional web data extraction infrastructures rely heavily on continuously provisioned virtual machines, "
        "monolithic scraping containers, or long-running cron servers that introduce significant administrative overhead, "
        "idle compute costs, fragile error recovery, and security vulnerabilities. This project presents The Cloud Scraper, "
        "a production-ready, fully serverless, event-driven data pipeline platform designed, deployed, and empirically evaluated "
        "for the Cloud Architecture and Design (BCSE355L) course."
    )
    add_body_paragraph(
        doc,
        "The system transitions traditional batch file ETL patterns into an autonomous, dual-mode web collection engine running "
        "live in the AWS Sydney region (ap-southeast-2). Orchestrated by AWS Step Functions across six independently scalable "
        "AWS Lambda microservices, the architecture supports both Automatic Generic HTML Extraction (heuristically discovering "
        "repeating DOM cards, titles, links, and text without requiring manual selector authoring) and Configurable CSS Selector "
        "Extraction. Ingestion is triggered via Amazon EventBridge on scheduled cron cadences or on-demand via Amazon API Gateway. "
        "Extracted payloads pass through an automated schema validation gate, undergo text sanitization and deterministic "
        "SHA-256 fingerprint deduplication, and are persisted across a dual-tier storage architecture: an immutable raw HTML "
        "data lake in Amazon S3 (SSE-AES256 encrypted) and an operational queryable store in Amazon DynamoDB. Fault tolerance "
        "is guaranteed through Step Functions fail-safe catch blocks routing persistent errors to an Amazon SQS Dead-Letter Queue "
        "(DLQ) with automated Amazon SNS notifications, ensuring zero pipeline stalling. The system is managed through an interactive "
        "web dashboard distributed globally via Amazon CloudFront (https://d3k1ttzl3bzy2k.cloudfront.net), enabling dynamic target "
        "registration, real-time Step Functions status polling, and client-side RFC 4180 CSV / JSON dataset export. Live cloud "
        "testing verified 100% pass rates across all ten formal acceptance criteria and forty-nine automated unit and integration tests, "
        "achieving an estimated per-run operational cost of $0.000165 USD."
    )

    add_body_paragraph(
        doc,
        "Serverless Computing, AWS Step Functions, AWS Lambda, Event-Driven ETL, Web Scraping, "
        "Data Hygiene, Deterministic Deduplication, Amazon DynamoDB, Amazon S3 Data Lake, Dead-Letter Queue, FinOps.",
        bold_prefix="Keywords: "
    )

    add_custom_heading(doc, "Acknowledgements", level=1)
    add_body_paragraph(
        doc,
        "We express our sincere gratitude to our course instructor, Prof. Venkatesh P, for his invaluable guidance, "
        "insightful lectures, and continuous academic encouragement throughout the Cloud Architecture and Design (BCSE355L) course. "
        "His rigorous architectural critiques and emphasis on cloud-native best practices, fault isolation, and cost awareness "
        "provided the foundational direction required to evolve this project from an initial concept into a fully deployable, "
        "serverless web intelligence platform."
    )
    add_body_paragraph(
        doc,
        "We also extend our appreciation to the School of Computer Science Engineering and Information Systems (SCOPE) at "
        "Vellore Institute of Technology for providing the technical infrastructure, laboratory facilities, and academic "
        "environment conducive to advanced applied engineering research. Finally, we thank our peers for their collaborative "
        "discussions on distributed systems, serverless patterns, and cloud security."
    )

    doc.add_page_break()


def create_table_of_contents_and_lists(doc):
    """Create explicit Table of Contents, List of Figures, and List of Tables."""
    add_custom_heading(doc, "Table of Contents", level=1)
    
    toc_items = [
        ("Abstract", "ii"),
        ("Acknowledgements", "iii"),
        ("List of Figures", "v"),
        ("List of Tables", "vi"),
        ("Chapter 1: Introduction", "1"),
        ("    1.1 Background and Motivation", "1"),
        ("    1.2 Context of Cloud Architecture and Design", "2"),
        ("    1.3 Project Scope and Intended Use Cases", "2"),
        ("    1.4 Document Organization", "3"),
        ("Chapter 2: Problem Statement and Objectives", "4"),
        ("    2.1 Problem Definition", "4"),
        ("    2.2 Original Proposal Objectives (Review 1 Baseline)", "4"),
        ("    2.3 Expected Outcomes", "5"),
        ("    2.4 Objective-to-Implementation Mapping", "6"),
        ("Chapter 3: Existing System vs. Proposed System", "8"),
        ("    3.1 Limitations of Traditional Web Scraping Implementations", "8"),
        ("    3.2 The Serverless Paradigm Shift", "9"),
        ("    3.3 Comparative Architectural Analysis", "9"),
        ("    3.4 Core Architectural Innovations", "11"),
        ("Chapter 4: Requirements Analysis", "13"),
        ("    4.1 Functional Requirements", "13"),
        ("    4.2 Non-Functional Requirements", "14"),
        ("    4.3 Software, Cloud, and Infrastructure Requirements", "15"),
        ("    4.4 External System Dependencies & Ethical Constraints", "16"),
        ("Chapter 5: System Architecture and Design", "17"),
        ("    5.1 Overall System Architecture", "17"),
        ("    5.2 Component Responsibilities and Layered Design", "18"),
        ("    5.3 End-to-End Serverless ETL Workflow", "20"),
        ("    5.4 Workflow Orchestration with AWS Step Functions", "21"),
        ("    5.5 Dual-Tier Cloud Storage Strategy", "23"),
        ("    5.6 Infrastructure as Code Design (AWS SAM)", "24"),
        ("Chapter 6: Implementation Details", "26"),
        ("    6.1 Public Web Dashboard & CloudFront CDN Distribution", "26"),
        ("    6.2 Job Dispatcher REST API Controller", "27"),
        ("    6.3 Dynamic Target Registration & Management", "28"),
        ("    6.4 Web Scraper Microservice (Extraction Engine)", "29"),
        ("    6.5 Schema Validation & Quality Gate Microservice", "31"),
        ("    6.6 Transformation & Deterministic Deduplication Microservice", "33"),
        ("    6.7 Dual-Tier Storage Microservice", "34"),
        ("    6.8 Fault Tolerance & Dead-Letter Queue Quarantine", "35"),
        ("    6.9 FinOps Telemetry & Alerting Microservice", "36"),
        ("Chapter 7: Technology Stack", "38"),
        ("    7.1 Architectural Technology Mapping", "38"),
        ("    7.2 Technical Rationale for Selected Cloud Services", "40"),
        ("Chapter 8: Testing and Validation", "42"),
        ("    8.1 Testing Methodology & Verification Framework", "42"),
        ("    8.2 Automated Unit & Integration Testing Suite", "42"),
        ("    8.3 Formal Acceptance Testing (Tests A through J)", "44"),
        ("    8.4 Fault Injection and Resilience Testing", "48"),
        ("Chapter 9: Results and Discussion", "49"),
        ("    9.1 Empirical Cloud Execution Metrics", "49"),
        ("    9.2 Cloud Storage Verification & Data Lineage Proofs", "50"),
        ("    9.3 Evaluation of Generic vs. Custom Selector Extraction", "51"),
        ("    9.4 Live Step Functions Execution Telemetry", "52"),
        ("    9.5 Synthesis Against Course Objectives", "54"),
        ("Chapter 10: Security, Reliability, and Cost Analysis", "55"),
        ("    10.1 Multi-Layer SSRF & Safety Protection Architecture", "55"),
        ("    10.2 Ethical Scraping, Rate Limiting & robots.txt Compliance", "57"),
        ("    10.3 Identity & Access Management (IAM) Least-Privilege Design", "58"),
        ("    10.4 Cryptographic Protection & Encryption at Rest", "59"),
        ("    10.5 High Availability, Fault Isolation & DLQ Quarantine", "59"),
        ("    10.6 FinOps Operational Cost Analysis", "60"),
        ("Chapter 11: Limitations and Future Enhancements", "62"),
        ("    11.1 Technical Limitations of Current Implementation", "62"),
        ("    11.2 Architectural Roadmap for Future Enhancements", "63"),
        ("Chapter 12: Conclusion", "65"),
        ("    12.1 Summary of Contributions", "65"),
        ("    12.2 Alignment with Cloud Architecture Design Principles", "66"),
        ("    12.3 Concluding Remarks", "66"),
        ("References", "67"),
        ("Appendices", "69"),
        ("    Appendix A: Representative API Request & Response Payloads", "69"),
        ("    Appendix B: Sample Structured Output Records", "71"),
        ("    Appendix C: Automated Pytest Test Suite Results Log", "73"),
        ("    Appendix D: AWS CloudFormation Resource Inventory", "75"),
    ]

    toc_table = doc.add_table(rows=len(toc_items), cols=2)
    toc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, (title, page) in enumerate(toc_items):
        row = toc_table.rows[idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width, c1.width = Inches(5.8), Inches(0.7)
        set_cell_background(c0, "FFFFFF")
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c0, top=20, bottom=20, left=40, right=40)
        set_cell_margins(c1, top=20, bottom=20, left=40, right=40)

        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before, p0.paragraph_format.space_after = Pt(0), Pt(0)
        r0 = p0.add_run(title)
        r0.font.name = "Calibri"
        r0.font.size = Pt(9.5)
        if title.startswith("Chapter") or title in ("Abstract", "Acknowledgements", "References", "Appendices"):
            r0.font.bold = True
            r0.font.color.rgb = COLOR_NAVY
        else:
            r0.font.color.rgb = COLOR_DARK

        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p1.paragraph_format.space_before, p1.paragraph_format.space_after = Pt(0), Pt(0)
        r1 = p1.add_run(page)
        r1.font.name = "Calibri"
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = COLOR_SLATE

    set_table_borders(toc_table, "FFFFFF")
    doc.add_page_break()

    # List of Figures
    add_custom_heading(doc, "List of Figures", level=1)
    figures_list = [
        ("Figure 5.1", "Overall AWS Serverless System Architecture (Sydney Region ap-southeast-2)", "18"),
        ("Figure 5.2", "Detailed Multi-Stage Serverless ETL Pipeline Execution Lifecycle", "20"),
        ("Figure 5.3", "AWS Step Functions State Machine Execution Flow with DLQ Routing", "22"),
        ("Figure 6.1", "Sequence Diagram for Dynamic Target Registration and Live Execution", "28"),
        ("Figure 6.2", "Dual-Tier Storage Architecture: S3 Data Lake vs. Amazon DynamoDB", "34"),
        ("Figure 9.1", "AWS Step Functions Management Console: State Machine Graph Orchestration", "52"),
        ("Figure 9.2", "AWS Step Functions Management Console: Successful Execution Telemetry", "53"),
        ("Figure 9.3", "AWS Step Functions Management Console: Failed Execution and DLQ Routing", "53"),
        ("Figure 10.1", "Multi-Layer SSRF & Safety Protection Architecture Flowchart", "56"),
    ]
    fig_table = doc.add_table(rows=len(figures_list), cols=2)
    fig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, (fnum, fdesc, fpage) in enumerate(figures_list):
        row = fig_table.rows[idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width, c1.width = Inches(5.8), Inches(0.7)
        set_cell_background(c0, "FFFFFF")
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c0, top=30, bottom=30, left=40, right=40)
        set_cell_margins(c1, top=30, bottom=30, left=40, right=40)

        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before, p0.paragraph_format.space_after = Pt(0), Pt(0)
        r0 = p0.add_run(f"{fnum}: {fdesc}")
        r0.font.name = "Calibri"
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = COLOR_DARK

        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p1.paragraph_format.space_before, p1.paragraph_format.space_after = Pt(0), Pt(0)
        r1 = p1.add_run(fpage)
        r1.font.name = "Calibri"
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = COLOR_SLATE

    set_table_borders(fig_table, "FFFFFF")

    # List of Tables
    add_custom_heading(doc, "List of Tables", level=1)
    tables_list = [
        ("Table 2.1", "Mapping of Review 1 Baseline Objectives to Live Cloud Implementation", "6"),
        ("Table 3.1", "Comparative Analysis: Traditional VM/Server Scraping vs. The Cloud Scraper", "10"),
        ("Table 7.1", "Comprehensive Technology Stack, Architectural Roles, and Implementation Proofs", "38"),
        ("Table 8.1", "Automated Pytest Suite Verification Summary (49 Tests Passed)", "43"),
        ("Table 8.2", "Formal Acceptance Test Results Matrix (Tests A through J)", "44"),
        ("Table 9.1", "Empirical Execution Outcomes and Performance Metrics Across Tested Targets", "49"),
        ("Table 10.1", "FinOps Operational Cost Model for 10,000 Monthly Invocations on AWS", "60"),
    ]
    tbl_table = doc.add_table(rows=len(tables_list), cols=2)
    tbl_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, (tnum, tdesc, tpage) in enumerate(tables_list):
        row = tbl_table.rows[idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width, c1.width = Inches(5.8), Inches(0.7)
        set_cell_background(c0, "FFFFFF")
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c0, top=30, bottom=30, left=40, right=40)
        set_cell_margins(c1, top=30, bottom=30, left=40, right=40)

        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before, p0.paragraph_format.space_after = Pt(0), Pt(0)
        r0 = p0.add_run(f"{tnum}: {tdesc}")
        r0.font.name = "Calibri"
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = COLOR_DARK

        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p1.paragraph_format.space_before, p1.paragraph_format.space_after = Pt(0), Pt(0)
        r1 = p1.add_run(tpage)
        r1.font.name = "Calibri"
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = COLOR_SLATE

    set_table_borders(tbl_table, "FFFFFF")
    doc.add_page_break()


print("Cover page, Abstract, Acknowledgement, and Lists generators defined.")
