"""
Build the complete, professionally formatted Microsoft Word (.docx) final project report
for The Cloud Scraper — Serverless Automated Data Pipeline.
Course: Cloud Architecture and Design (BCSE355L)
"""

import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

PROJECT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUTPUT_DOCX = os.path.join(PROJECT_DIR, "The_Cloud_Scraper_Final_Project_Report.docx")
FIGURES_DIR = os.path.join(PROJECT_DIR, "figures")
IMAGES_DIR = os.path.join(PROJECT_DIR, "images")

# Palette constants
COLOR_NAVY = RGBColor(15, 23, 42)      # #0f172a
COLOR_BLUE = RGBColor(2, 132, 199)     # #0284c7
COLOR_SLATE = RGBColor(71, 85, 105)    # #475569
COLOR_DARK = RGBColor(30, 41, 59)      # #1e293b
COLOR_MUTED = RGBColor(100, 116, 139)  # #64748b

HEX_NAVY = "0F172A"
HEX_BLUE = "0284C7"
HEX_LIGHT_BG = "F8FAFC"
HEX_ALT_ROW = "F1F5F9"
HEX_BORDER = "CBD5E1"
HEX_CALLOUT_BG = "F0F9FF"
HEX_CODE_BG = "F1F5F9"


def set_cell_background(cell, hex_color):
    """Set the background color of a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set padding/margins for a table cell in dxa (1 pt = 20 dxa)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)


def set_table_borders(table, border_color=HEX_BORDER):
    """Apply clean light borders to a table."""
    tbl_pr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tbl_pr.append(borders)


def format_table_header_and_rows(table, col_widths, alignments=None):
    """Style table with Navy header, white text, and alternating rows."""
    set_table_borders(table)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for row_idx, row in enumerate(table.rows):
        # Prevent row split across pages
        tr_pr = row._tr.get_or_add_trPr()
        tr_pr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

        # Header row configuration
        if row_idx == 0:
            tr_pr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

        for col_idx, cell in enumerate(row.cells):
            cell.width = Inches(col_widths[col_idx])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

            if row_idx == 0:
                set_cell_background(cell, HEX_NAVY)
                set_cell_margins(cell, top=140, bottom=140, left=140, right=140)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    if alignments and col_idx < len(alignments):
                        p.alignment = alignments[col_idx]
                    for run in p.runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
                        run.font.size = Pt(9.5)
                        run.font.name = "Calibri"
            else:
                bg = HEX_ALT_ROW if row_idx % 2 == 1 else "FFFFFF"
                set_cell_background(cell, bg)
                set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    if alignments and col_idx < len(alignments):
                        p.alignment = alignments[col_idx]
                    for run in p.runs:
                        run.font.size = Pt(9.0)
                        run.font.color.rgb = COLOR_DARK
                        run.font.name = "Calibri"


def add_custom_heading(doc, text, level):
    """Add a styled heading with controlled spacing and page binding."""
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.bold = True

    if level == 1:
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        run.font.size = Pt(17)
        run.font.color.rgb = COLOR_NAVY
    elif level == 2:
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        run.font.size = Pt(13.5)
        run.font.color.rgb = COLOR_BLUE
    elif level == 3:
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        run.font.size = Pt(11.5)
        run.font.color.rgb = COLOR_SLATE
    return p


def add_body_paragraph(doc, text, bold_prefix=None, space_after=5):
    """Add a standard formatted body paragraph."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15

    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Calibri"
        r_pre.font.bold = True
        r_pre.font.size = Pt(11)
        r_pre.font.color.rgb = COLOR_NAVY

    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(11)
    r.font.color.rgb = COLOR_DARK
    return p


def add_bullet_item(doc, text, bold_prefix=None):
    """Add a clean bullet item."""
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15

    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Calibri"
        r_pre.font.bold = True
        r_pre.font.size = Pt(10.5)
        r_pre.font.color.rgb = COLOR_NAVY

    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(10.5)
    r.font.color.rgb = COLOR_DARK
    return p


def add_code_block(doc, code_text):
    """Add a shaded code block with a thin left accent border."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, HEX_CODE_BG)
    set_cell_margins(cell, top=120, bottom=120, left=160, right=160)

    tc_pr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="18" w:space="0" w:color="{HEX_BLUE}"/>'
        f'<w:top w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tc_pr.append(borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(code_text)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = COLOR_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_callout(doc, title, text):
    """Add a shaded callout box for architectural notes and key takeaways."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, HEX_CALLOUT_BG)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)

    tc_pr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{HEX_BLUE}"/>'
        f'<w:top w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tc_pr.append(borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    r_title = p.add_run(title)
    r_title.font.name = "Calibri"
    r_title.font.bold = True
    r_title.font.size = Pt(10.5)
    r_title.font.color.rgb = COLOR_BLUE

    p2 = cell.add_paragraph()
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.15
    r_text = p2.add_run(text)
    r_text.font.name = "Calibri"
    r_text.font.size = Pt(10)
    r_text.font.color.rgb = COLOR_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_figure_with_caption(doc, image_path, caption_text, width_inches=6.2):
    """Insert a centered figure with a formal caption underneath."""
    if not os.path.exists(image_path):
        print(f"Warning: Image {image_path} not found. Skipping figure.")
        return

    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(3)
    p_img.paragraph_format.keep_with_next = True
    run_img = p_img.add_run()
    run_img.add_picture(image_path, width=Inches(width_inches))

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(0)
    p_cap.paragraph_format.space_after = Pt(10)
    p_cap.paragraph_format.keep_with_next = False
    run_cap = p_cap.add_run(caption_text)
    run_cap.font.name = "Calibri"
    run_cap.font.size = Pt(9.5)
    run_cap.font.italic = True
    run_cap.font.color.rgb = COLOR_SLATE


def add_page_number_to_run(run):
    """Add a dynamic Word page number XML field to a run."""
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'PAGE'
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)


def add_numpages_to_run(run):
    """Add a dynamic Word NUMPAGES XML field to a run."""
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'NUMPAGES'
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)


def setup_document_styles_and_margins(doc):
    """Configure A4 margins, headers, footers, and default typography."""
    for s in doc.sections:
        s.page_width = Inches(8.27)
        s.page_height = Inches(11.69)
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)
        s.header_distance = Inches(0.5)
        s.footer_distance = Inches(0.5)

    # Base style
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = COLOR_DARK


print("Setup helper functions initialized.")
